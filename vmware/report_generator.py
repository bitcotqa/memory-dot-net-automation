"""
Excel (+ CSV) report generation: main per-record sheet and a summary sheet.
"""

import csv
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

MAIN_COLUMNS = [
    "S.No",
    "Category",
    "Part URL",
    "Part Number",
    "Expected DWPD",
    "Actual DWPD",
    "DWPD Result",
    "Expected Form Factor",
    "Actual Form Factor",
    "Form Factor Result",
    "Expected Capacity",
    "Actual Capacity",
    "Capacity Result",
    "Expected MFR Part No",
    "Actual MFR Part No",
    "MFR Part No Result",
    "Expected Part Description",
    "Actual Part Description",
    "Part Description Result",
    "Expected Interface",
    "Actual Interface",
    "Interface Result",
    "Expected Part Specification",
    "Actual Part Specification",
    "Part Specification Result",
    "Page Status",
    "Overall Result",
    "Comments",
]

FIELD_RESULT_COLUMNS = {
    "Part Number": None,  # part number itself is validated but has no dedicated
                           # Expected/Actual pair in the fixed column list above;
                           # its comparison result is folded into Overall Result
                           # and Comments per the requested column layout.
    "DWPD": "DWPD Result",
    "Form Factor": "Form Factor Result",
    "Capacity": "Capacity Result",
    "MFR Part No": "MFR Part No Result",
    "Part Description": "Part Description Result",
    "Interface": "Interface Result",
    "Part Specification": "Part Specification Result",
}

RESULT_FILL = {
    "MATCHED": "C6EFCE",
    "PARTIAL MATCH": "FFEB9C",
    "UNMATCHED": "FFC7CE",
    "NOT FOUND": "FFC7CE",
    "NOT APPLICABLE": "D9D9D9",
    "BLOCKED": "D9D9D9",
    "ERROR": "FFC7CE",
}


def write_csv_report(results, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=MAIN_COLUMNS)
        writer.writeheader()
        for r in results:
            writer.writerow({k: r.get(k, "") for k in MAIN_COLUMNS})


# ---------------------------------------------------------------------------
# Duplicate-row + Part URL validation report (row_level_validator.py).
# Written as a separate CSV, independent of the browser-scrape report above,
# since it covers the entire source CSV rather than a sampled subset.
# ---------------------------------------------------------------------------

DUPLICATE_URL_COLUMNS = [
    "csv_row_index",
    "Category",
    "Part URL",
    "Duplicate Row",
    "Duplicate Group Id",
    "Duplicate Of Row(s)",
    "Part URL Validation Result",
    "Part URL Validation Comments",
]


def write_duplicate_url_csv_report(results, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=DUPLICATE_URL_COLUMNS)
        writer.writeheader()
        for r in results:
            writer.writerow({k: r.get(k, "") for k in DUPLICATE_URL_COLUMNS})


def compute_summary(results):
    total = len(results)
    ssd = sum(1 for r in results if r["Category"].strip().lower() == "ssd")
    hdd = sum(1 for r in results if r["Category"].strip().lower() == "hdd")

    overall_counter = Counter(r["Overall Result"] for r in results)

    field_names = ["Part Number", "DWPD", "Form Factor", "Capacity",
                    "MFR Part No", "Part Description", "Interface",
                    "Part Specification"]
    field_stats = {}
    for field in field_names:
        result_col = FIELD_RESULT_COLUMNS.get(field)
        if result_col is None:
            counts = Counter(r.get("_part_number_result", "") for r in results)
        else:
            counts = Counter(r.get(result_col, "") for r in results)
        field_stats[field] = {
            "MATCHED": counts.get("MATCHED", 0),
            "UNMATCHED": counts.get("UNMATCHED", 0),
            "NOT FOUND": counts.get("NOT FOUND", 0),
            "NOT APPLICABLE": counts.get("NOT APPLICABLE", 0),
            "PARTIAL MATCH": counts.get("PARTIAL MATCH", 0),
        }

    return {
        "total": total,
        "ssd": ssd,
        "hdd": hdd,
        "overall": overall_counter,
        "field_stats": field_stats,
    }


def write_excel_report(results, path):
    wb = Workbook()

    ws = wb.active
    ws.title = "Results"
    ws.append(MAIN_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")

    for r in results:
        row = [r.get(c, "") for c in MAIN_COLUMNS]
        ws.append(row)

    for row_idx in range(2, ws.max_row + 1):
        for col_name in ("DWPD Result", "Form Factor Result", "Capacity Result",
                          "MFR Part No Result", "Part Description Result",
                          "Interface Result", "Part Specification Result",
                          "Overall Result"):
            col_idx = MAIN_COLUMNS.index(col_name) + 1
            cell = ws.cell(row=row_idx, column=col_idx)
            fill_color = RESULT_FILL.get(str(cell.value).strip())
            if fill_color:
                cell.fill = PatternFill("solid", fgColor=fill_color)

    for col_idx, col_name in enumerate(MAIN_COLUMNS, start=1):
        width = 18
        if col_name in ("Part URL", "Comments", "Actual Part Specification",
                        "Expected Part Specification"):
            width = 45
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width
    ws.freeze_panes = "A2"

    summary = compute_summary(results)
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Metric", "Value"])
    ws2["A1"].font = Font(bold=True)
    ws2["B1"].font = Font(bold=True)

    ws2.append(["Total records tested", summary["total"]])
    ws2.append(["SSD records tested", summary["ssd"]])
    ws2.append(["HDD records tested", summary["hdd"]])
    for key in ("MATCHED", "PARTIAL MATCH", "UNMATCHED", "NOT FOUND", "BLOCKED", "ERROR"):
        ws2.append([key, summary["overall"].get(key, 0)])

    ws2.append([])
    ws2.append(["Field", "Matched", "Unmatched", "Not Found", "Not Applicable", "Partial Match"])
    for cell in ws2[ws2.max_row]:
        cell.font = Font(bold=True)
    for field, stats in summary["field_stats"].items():
        ws2.append([
            field,
            stats["MATCHED"],
            stats["UNMATCHED"],
            stats["NOT FOUND"],
            stats["NOT APPLICABLE"],
            stats["PARTIAL MATCH"],
        ])

    for col_letter in ("A", "B", "C", "D", "E", "F"):
        ws2.column_dimensions[col_letter].width = 22

    wb.save(path)
    return summary


# ---------------------------------------------------------------------------
# Generic styled Excel report - usable by ANY checker script (not tied to
# MAIN_COLUMNS like write_excel_report above). Header cells are coloured by
# column group so Expected / Actual / Result columns are easy to tell apart,
# and result cells reuse the same RESULT_FILL colours as the main report.
# ---------------------------------------------------------------------------

# Header colour per column group (dark fills, paired with white bold text).
HEADER_FILL_EXPECTED = "1F4E78"  # dark blue   - "Expected ..." (value from CSV)
HEADER_FILL_ACTUAL = "2E75B6"    # medium blue - "Actual ..." (value compared against)
HEADER_FILL_RESULT = "548235"    # green       - "... Result" (comparison outcome)
HEADER_FILL_OTHER = "404040"     # dark grey   - everything else (ids, URL, comments)

# Columns holding long text get a wider column; all others use the default.
WIDE_COLUMN_WIDTH = 45
DEFAULT_COLUMN_WIDTH = 20


def header_fill_for(column_name):
    """Pick the header colour for a column based on its name prefix/suffix."""
    if column_name.startswith("Expected"):   # CSV-side value columns
        return HEADER_FILL_EXPECTED
    if column_name.startswith("Actual"):     # compared-against value columns
        return HEADER_FILL_ACTUAL
    if column_name.endswith("Result"):       # MATCHED / UNMATCHED / ... columns
        return HEADER_FILL_RESULT
    return HEADER_FILL_OTHER                 # identifiers, URL, comments, etc.


def write_styled_excel_report(results, columns, path, sheet_title="Results",
                              wide_columns=(), freeze_panes="A2"):
    """Write result dicts to an .xlsx with coloured headers and result cells.

    results      - list of dicts, one per report row (same dicts used for CSV)
    columns      - ordered column names to write (also the header row)
    path         - output .xlsx file path
    sheet_title  - name of the worksheet tab
    wide_columns - column names that hold long text and need a wider column
    freeze_panes - top-left cell that stays unfrozen (e.g. "C2" keeps the
                   header row and the first two columns visible on scroll)
    """
    wb = Workbook()                            # new, empty workbook
    ws = wb.active                             # its default (first) sheet
    ws.title = sheet_title                     # rename the sheet tab

    ws.append(list(columns))                   # row 1: the header names
    for r in results:                          # rows 2..N: one per result dict
        ws.append([r.get(c, "") for c in columns])  # missing keys -> blank cell

    header_font = Font(bold=True, color="FFFFFF")   # white bold header text
    header_align = Alignment(wrap_text=True, horizontal="center", vertical="center")
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx)  # header cell of this column
        cell.fill = PatternFill("solid", fgColor=header_fill_for(col_name))  # group colour
        cell.font = header_font                # readable on the dark fill
        cell.alignment = header_align          # wrap long names, centre them
        width = WIDE_COLUMN_WIDTH if col_name in wide_columns else DEFAULT_COLUMN_WIDTH
        ws.column_dimensions[cell.column_letter].width = width  # set column width
    ws.row_dimensions[1].height = 45           # taller header row for wrapped names

    result_col_indexes = [i for i, c in enumerate(columns, start=1) if c.endswith("Result")]
    for row_idx in range(2, ws.max_row + 1):   # every data row (skip header)
        for col_idx in result_col_indexes:     # only the "... Result" columns
            cell = ws.cell(row=row_idx, column=col_idx)
            fill_color = RESULT_FILL.get(str(cell.value).strip())  # reuse main report colours
            if fill_color:                     # unknown values stay unfilled
                cell.fill = PatternFill("solid", fgColor=fill_color)

    ws.freeze_panes = freeze_panes             # keep header (and id columns) on screen
    ws.auto_filter.ref = ws.dimensions         # filter dropdown on every header cell
    wb.save(path)                              # write the .xlsx to disk
